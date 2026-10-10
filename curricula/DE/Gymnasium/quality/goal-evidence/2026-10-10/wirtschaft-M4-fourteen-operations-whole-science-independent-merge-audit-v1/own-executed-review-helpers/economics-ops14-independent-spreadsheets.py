from pathlib import Path
from fractions import Fraction as F
import openpyxl,xml.etree.ElementTree as ET,zipfile,re,json,hashlib,ast
from openpyxl.chart import BarChart,Reference
R=Path('/home/enpasos/projects/skillpilot');O=R/Path('/tmp/economics-ops14-independent-own-path.txt').read_text().strip();D=R/'curricula/DE/Gymnasium/quality/goal-evidence/2026-10-10/wirtschaft-M4-fourteen-finance-business-and-representation-KEEP-first-current597-author-b-v1'
def h(p):return hashlib.sha256(p.read_bytes()).hexdigest()
wb=openpyxl.Workbook();wb.remove(wb.active)
def sheet(name,rows,note,charts):
 s=wb.create_sheet(name)
 for row in rows:s.append(row)
 s.cell(len(rows)+3,1,note)
 s.freeze_panes='B2';s.column_dimensions['A'].width=35
 for c in range(2,7):s.column_dimensions[openpyxl.utils.get_column_letter(c)].width=24
 for col,title,ylabel,place in charts:
  ch=BarChart();ch.title=title;ch.y_axis.title=ylabel;ch.x_axis.title='Abgrenzung laut Tabelle';ch.y_axis.scaling.min=0
  ch.add_data(Reference(s,min_col=col,max_col=col,min_row=1,max_row=len(rows)),titles_from_data=True);ch.set_categories(Reference(s,min_col=1,min_row=2,max_row=len(rows)));s.add_chart(ch,place)
 return s
sheet('8328_A_2024',[['WZ08','Statistische Unternehmen','Beschäftigungsverhältnisse','Beschäftigung je Unternehmen'],['C',198362,7801230,'=C2/B2'],['G',538248,6017332,'=C3/B3']],'Amtlicher Auszug 2024, Stand 3.8.2026. Beschäftigungsverhältnisse, keine eindeutigen Personenzahlen. Modelländerung ist getrennt fiktiv.',[(2,'Unternehmen C/G 2024','Anzahl statistischer Unternehmen','G2'),(3,'Abhängige Beschäftigung C/G 2024','Jahresdurchschnitt Beschäftigungsverhältnisse','G18')])
s=sheet('8328_B_Handwerk',[['WZ08','Rechtliche Einheiten in Tausend','Handwerk in Tausend','Anteil aus gerundeten Zahlen'],['C',210,86,'=C2/B2'],['F',388,253,'=C3/B3']],'Amtlicher Auszug 2024, Stand 23.4.2026. Originalanteile40.9/65.2%, gerechnete gerundete Zahlen abweichend. Keine Gleichsetzung mit statistischen Unternehmen in A.',[(4,'Handwerkanteil aus gerundeten Zahlen','Anteil (0 bis1)','G2')]);s['D2'].number_format=s['D3'].number_format='0.00%'
s=sheet('685_A_GuV',[['Jahr','Umsatz in Tausend EUR','Gewinn in Tausend EUR','Gewinn/Umsatz'],[2023,100,12,'=C2/B2'],[2024,130,13,'=C3/B3'],[2025,150,9,'=C4/B4']],'Fiktive gleich abgegrenzte Jahre. Unterschiedliche Größen getrennt gezeichnet, keine zweite versteckte Achse. Kein Kausal-, Liquiditäts- oder Sicherheitsbeweis.',[(2,'Umsatzentwicklung','Tausend EUR','G2'),(3,'Gewinnentwicklung','Tausend EUR','G18')]);
for c in ['D2','D3','D4']:s[c].number_format='0.00%'
s=sheet('685_B_Bilanz',[['Fiktiver Stichtag','Eigenkapital in Tausend EUR','Schulden in Tausend EUR','Summe','EK/Summe','Schulden/Summe'],['31.12.',90,210,'=B2+C2','=B2/D2','=C2/D2']],'Aktiva insgesamt300 sind die Gegenseite, kein dritter Finanzierungsanteil. Lieferanten brauchen Fälligkeiten, Zahlungsmittel und Liquiditätsnachweise.',[(2,'Finanzierungsstruktur am Stichtag','Tausend EUR','H2')]);ch=s._charts[0];ch.add_data(Reference(s,min_col=3,max_col=3,min_row=1,max_row=2),titles_from_data=True)
for c in ['E2','F2']:s[c].number_format='0.00%'
before=O/'independent-four-case-six-chart-actual-work.before.xlsx';wb.save(before)
mutations=[('8328_A_2024','C3',6117332),('8328_B_Handwerk','C3',250),('685_A_GuV','C4',18),('685_B_Bilanz','B2',110),('685_B_Bilanz','C2',190)]
for sh,c,v in mutations:wb[sh][c]=v
wb['8328_A_2024']['A8']='Fiktive Änderung nur Beschäftigung Handel +100000; keine amtliche Revision.'
wb['8328_B_Handwerk']['A8']='Fiktive Änderung Bauhandwerk253→250 bei unverändertem388; keine amtliche Revision.'
wb['685_A_GuV']['A9']='Fiktive Gewinnänderung2025 9→18 bei unverändertemUmsatz150.'
wb['685_B_Bilanz']['A7']='Fiktive Finanzierung bei Summe300: Eigenkapital110, Schulden190.'
after=O/'independent-four-case-six-chart-actual-work.after.xlsx';wb.save(after)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
def evaluate(s,c,seen=()):
 if c in seen:raise ValueError('cycle')
 v=s[c].value
 if isinstance(v,(int,float)):return F(str(v))
 if not isinstance(v,str) or not v.startswith('='):raise ValueError((s.title,c,v))
 ex=re.sub(r'\b([A-Z]+[0-9]+)\b',lambda m:'('+str(evaluate(s,m[1],seen+(c,)))+')',v[1:]); t=ast.parse(ex,mode='eval')
 def e(n):
  if isinstance(n,ast.Expression):return e(n.body)
  if isinstance(n,ast.Constant):return F(str(n.value))
  if isinstance(n,ast.BinOp):
   a,b=e(n.left),e(n.right)
   if isinstance(n.op,ast.Add):return a+b
   if isinstance(n.op,ast.Sub):return a-b
   if isinstance(n.op,ast.Mult):return a*b
   if isinstance(n.op,ast.Div):return a/b
  raise ValueError(ast.dump(n))
 return e(t)
expected0={'8328_A_2024':{'D2':F(7801230,198362),'D3':F(6017332,538248)},'8328_B_Handwerk':{'D2':F(86,210),'D3':F(253,388)},'685_A_GuV':{'D2':F(12,100),'D3':F(13,130),'D4':F(9,150)},'685_B_Bilanz':{'D2':F(300),'E2':F(90,300),'F2':F(210,300)}}
expected1=json.loads(json.dumps({s:{c:str(v) for c,v in cs.items()} for s,cs in expected0.items()}));expected1={s:{c:F(v) for c,v in cs.items()} for s,cs in expected1.items()}
expected1['8328_A_2024']['D3']=F(6117332,538248);expected1['8328_B_Handwerk']['D3']=F(250,388);expected1['685_A_GuV']['D4']=F(18,150);expected1['685_B_Bilanz']['E2']=F(110,300);expected1['685_B_Bilanz']['F2']=F(190,300)
checks=[];frames=[]
for origin,prefix in [('independent',O),('author-input',D)]:
 paths=[before,after] if origin=='independent' else [D/'own-four-spreadsheet-cases.before-actual-input-change.xlsx',D/'own-four-spreadsheet-cases.after-actual-input-change.xlsx']
 for no,p in enumerate(paths):
  w=openpyxl.load_workbook(p,data_only=False);assert len(w.sheetnames)==4;z=zipfile.ZipFile(p);chartRows=[]
  for sh,cs in (expected0 if no==0 else expected1).items():
   for c,v in cs.items():
    actual=evaluate(w[sh],c);checks.append({'origin':origin,'path':str(p.relative_to(R)),'sheet':sh,'cell':c,'savedFormula':w[sh][c].value,'actual':str(actual),'expected':str(v),'pass':actual==v});assert actual==v
  for n in z.namelist():
   if n.startswith('xl/charts/chart') and n.endswith('.xml'):
    r=ET.fromstring(z.read(n));refs=[e.text for e in r.findall('.//c:f',ns)];assert refs
    resolved=[]
    for ref in refs:
     sh,ran=ref.split('!');sh=sh.strip("'");ran=ran.replace('$','');cells=list(w[sh][ran]) if ':' in ran else [(w[sh][ran],)]
     values=[]
     for row in cells:
      for c in row:
       if c.data_type=='f':values.append(str(evaluate(w[sh],c.coordinate)))
       else:values.append(c.value)
     resolved.append({'reference':ref,'actualValues':values})
    chartRows.append({'member':n,'memberSHA':hashlib.sha256(z.read(n)).hexdigest(),'dataResolvedFromActualSavedCells':resolved})
  assert len(chartRows)==6
  frames.append({'origin':origin,'path':str(p.relative_to(R)),'sha256':h(p),'charts':chartRows,'formulaCount':sum(c.data_type=='f' for sh in w for row in sh for c in row)})
# A concrete edited denominator formula must be rejected by the same independent expected-value comparison.
w=openpyxl.load_workbook(before);w['8328_A_2024']['D2']='=B2/C2';neg=O/'independent-real-formula-inversion-negative.xlsx';w.save(neg);assert evaluate(w['8328_A_2024'],'D2')!=expected0['8328_A_2024']['D2']
result={'role':'independent actual XLSX archive/cell/formula/reference replay and own separate saved work; no graphical spreadsheet-host execution claimed','frames':frames,'checks':checks,'actualCheckCount':len(checks),'errors':0,'ownRealFiveInputMutations':mutations,'ownNegative':{'path':str(neg.relative_to(R)),'sha256':h(neg),'sheet':'8328_A_2024','cell':'D2','change':'C2/B2→B2/C2','actualRejected':True},'authorInputOriginalsUnchanged':True,'actualHostRenderingClaim':False,'actualHostRecalculationClaim':False}
(O/'actual-independent-four-XLSX-frames-forty-Fraction-checks-six-chart-references-and-real-formula-negative.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print('40 checks PASS; four frames/24 chart bindings; real formula inversion rejected')
