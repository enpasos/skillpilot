"""Independent E40 delta review. Temporary spreadsheet dependency stays outside repo."""
import json, hashlib, copy, sys, zipfile
from pathlib import Path
from datetime import datetime,timezone
import jsonschema
sys.path.insert(0, '/tmp/skillpilot-e40-independent-counterworkbook-deps')
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

HERE=Path(__file__).resolve().parent
REPO=next(p for p in HERE.parents if (p/'AGENTS.md').exists())
PREFIX=REPO/'curricula/DE/Gymnasium/quality/goal-evidence/2026-10-09'
AUTHOR=PREFIX/'wirtschaft-E40-bounded-independent-findings-author-successor-v2'
PRIOR=PREFIX/'wirtschaft-four-E40-environment-finance-production-team-independent-whole-material-review-v1'
def save(n,o): (HERE/n).write_text(json.dumps(o,ensure_ascii=False,indent=2)+'\n')
def objsha(o):return 'sha256:'+hashlib.sha256(json.dumps(o,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
old=json.loads((PRIOR/'whole-four-current-materials.only-two-independent-machine-material-statuses-released.inert.json').read_text())
new=json.loads((AUTHOR/'whole-two-finance-and-project-mandatory-execution-cap14.DRAFT-author-successors.json').read_text())
changes=[];judgments=[]
for g in new:
 before=next(x for x in old if x['id']==g['id'])
 assert {k:v for k,v in g.items() if k!='examData'}=={k:v for k,v in before.items() if k!='examData'}
 assert {k:v for k,v in g['examData'].items() if k not in ['taskContent','taskContentEn','solutionContent','solutionContentEn','scoring']}=={k:v for k,v in before['examData'].items() if k not in ['taskContent','taskContentEn','solutionContent','solutionContentEn','scoring']}
 fields=[]
 for field in ['taskContent','taskContentEn','solutionContent','solutionContentEn']:
  orig=before['examData'][field];cur=g['examData'][field]
  corrected=orig
  if g['id'].startswith('642') and field=='taskContent':
   corrected=orig.replace('Die Gesellschafterhaftung ist grundsätzlich auf Gesellschaftsvermögen begrenzt', 'Für die Verbindlichkeiten der GmbH haftet den Gesellschaftsgläubigern grundsätzlich nur das Gesellschaftsvermögen; die Gesellschafter haften für diese Gesellschaftsverbindlichkeiten grundsätzlich nicht persönlich')
  if g['id'].startswith('642') and field=='taskContentEn':
   corrected=orig.replace('Shareholder liability is normally limited to company assets','For the GmbH’s obligations creditors generally have recourse to company assets; shareholders generally are not personally liable for those company obligations')
  assert cur.startswith(corrected)
  addition=cur[len(corrected):]
  assert ('14BE' in addition) if not field.endswith('En') else ('14points' in addition)
  assert ('keine fehlerfrei' in addition) if not field.endswith('En') else ('no perfect6' in addition)
  fields.append({'field':'examData.'+field,'wholeOldFieldSHA256':objsha(orig),'wholeNewFieldSHA256':objsha(cur),'actualWholeAppendedCondition':addition,'originalWholeBodyExactAfterOnlySpecifiedActorCorrection':True})
 bef=before['examData']['scoring'];cur=g['examData']['scoring'];assert cur['maxPoints']==bef['maxPoints']==24 and cur['passingPoints']==bef['passingPoints']==15
 for i,(a,b) in enumerate(zip(bef['steps'],cur['steps'])):
  assert a['id']==b['id'] and a['points']==b['points']==6
  if a!=b:
   assert b['description'].startswith(a['description'])
   assert i==(3 if g['id'].startswith('642') else 2)
   fields.append({'field':'examData.scoring.steps['+str(i)+'].description','actualOnlyAppendedCondition':b['description'][len(a['description']):]})
 changes.append({'goalId':g['id'],'wholeOuterContractRequiresCoveredGoalIDsSourceMaterialNumbersExact':True,'exactBoundedFields':fields})
 judgments.append({'goalId':g['id'],'verdict':'KEEP','wholeCurrentAuthorObjectSHA256':objsha(g),
  'priorWholeSubstantiveReadReusedByExactness':True,'actualAllNewDEENConditionsAndActorChangesActuallyRead':True,
  'reason':('Company assets/company creditors distinguished from normally non-personally-liable shareholders and the stipulated separate valid guarantee; whole execution condition blocks the real20-point no-spreadsheet counteranswer.' if g['id'].startswith('642') else 'Already demanded project, team/explicit role and linked spreadsheet execution are each essential; wholecap14 blocks the conservative actual16-point no-spreadsheet answer and the additional16-point no-project/role case.'),
  'noLoweredQualityThreshold':'Same24total/15pass/4×6 allocations; no perfect score, extra task count, student agreement or new private data condition. Executed imperfect work retains normal partial credit; equivalent verifiable actual runs remain possible.',
  'humanReview':'pending','noSourceCourseDVPOrMasteryApproval':True})

priorcases=json.loads((PRIOR/'actual-eight-complete-independent-counteranswers-and-individual-rubric-marks.json').read_text())['cases']
rechecks=[]
for c in priorcases:
 if c['examGoalId'] not in {g['id'] for g in new}:continue
 raw=c['totalPoints'];cap=c['missingEssentialActualPerformance'] is not None
 score=min(raw,14) if cap else raw
 assert score<15
 rechecks.append({'caseId':c['id'],'wholeUnchangedPriorCounteranswerObjectSHA256':objsha(c),'oldActualStepScores':[s['awardedStepPoints'] for s in c['steps']],
  'oldActualRawTotal':raw,'newWholeExecutionConditionApplies':cap,'newTotal':score,'passingPoints':15,'newPass':False,'sourceOriginalWholeCounteranswerRetainedNotRewritten':True})

# Fresh extra case: actually supplied linked statistics spreadsheet, but project and
# role answers remain unexecuted plans. This is a reviewer fixture, not learner work.
wb=Workbook();ws=wb.active;ws.title='Data2024'
for row in [ ['Wirtschaftsabschnitt2024','Unternehmen Anzahl','Beschäftigungsverhältnisse Jahresdurchschnitt'],
 ['Verarbeitendes Gewerbe',198362,7801230],['Handel inklKfz-Reparatur',538248,6017332] ]:ws.append(row)
ws['E1']='Beziehungen jeUnternehmen';ws['E2']='=C2/B2';ws['E3']='=C3/B3'
ws['A6']='Destatis Unternehmensregister2024; Stand3August2026; SitzDE/WZ2008; Beziehungen keine einmalig gezähltenPersonen.'
ws['A7']='Reviewfixture only; actual chart/data XML created; no observed learner, externalteam or fullprojectexecution.'
for col,title,pos in [(2,'Unternehmen nachWirtschaftsabschnitt2024','G2'),(3,'Beschäftigungsverhältnisse Jahresdurchschnitt2024','G18')]:
 chart=BarChart();chart.title=title;chart.y_axis.title='Anzahl';chart.x_axis.title='Wirtschaftsabschnitt/WZ2008';chart.add_data(Reference(ws,min_col=col,min_row=1,max_row=3),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=3));ws.add_chart(chart,pos)
book=HERE/'actual-independent-no-project-or-role-execution.statistics-linked-chart-fixture.xlsx';wb.save(book)
loaded=load_workbook(book);assert len(loaded['Data2024']._charts)==2 and loaded['Data2024']['E2'].value=='=C2/B2'
with zipfile.ZipFile(book) as z:
 chartxml={n:z.read(n).decode() for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')}
 assert len(chartxml)==2 and any('$B$2:$B$3' in x for x in chartxml.values()) and any('$C$2:$C$3' in x for x in chartxml.values())
previous=next(x for x in priorcases if x['id']=='project-all-prose-with-no-executed-products')
extra={
 'id':'project-linked-spreadsheet-but-project-and-roles-never-executed','examGoalId':'b2003473-c85f-57e9-83be-9bbc0a922b32',
 'syntheticReviewerFixtureNotObservedLearnerOrAuthorCounterReuse':True,
 'completeFourTaskAnswersDe':[
  'Die Frage lautet, wie sich die Branchenstruktur2024 unterscheidet. Geplanter Ablauf: Quelle prüfen, rechnen, getrennte Charts, Bericht und Auswertung. Ich stelle diesen Projektplan lediglich vor und habe kein Projekt-/Entscheidungsberichtprodukt oder methodische Gesamtauswertung im Projektzyklus ausgeführt. Die separat beiliegende statistische Worksheetübung ist tatsächlich vorhanden, ersetzt aber das geplante Gesamtprojekt nicht. Amtliche Werte, eigene Quotienten und fiktive Rollenausfallbedingung sind getrennt.',
  'Vorgeschlagen sind Daten/Grafik/Redaktion. Quelle liegt vor Rechnung, diese vorCharts, Analyse vorBericht. Beim Ausfall würdeDaten die Grafikrolle übernehmen. Ich habe diese Rollen weder mitTeam noch in einer gekennzeichneten tatsächlichen Rollenarbeit umgesetzt und keinen ausgeführten Abstimmungs-, Änderungs- oder Statusdurchlauf dokumentiert. Eine Toolwahl könnteVersionsstände sichtbar machen; hier bleibt sie eineAbsicht.',
  'Die beigefügte tatsächliche Worksheetübung enthält Quellenjahr, getrennte zwei datenverknüpfte Balkendiagramme mit2024/Anzahl/Branchen sowie198362/7801230 und538248/6017332. FormelnC2/B2 undC3/B3 entsprechen39.328 und11.179 unabhängig geprüfterRundung. Handelsfirmen sindmehr, Gewerbe-Beschäftigungsverhältnisse mehr. Beziehungen sindJahresdurchschnittswerte, Mehrfachjobskeine einzigartigenMenschen undMittelnichtMedian. Eine weitereErkundung sollte örtlicheTätigkeiten/Qualifikation/Angebote prüfen, keineindividuelleChancegarantie.',
  previous['steps'][3]['completeAnswerDe']],
 'actualCounterArtefact':{'path':book.relative_to(REPO).as_posix(),'sha256':'sha256:'+hashlib.sha256(book.read_bytes()).hexdigest(),'linkedChartCount':2,'formulas':['=C2/B2','=C3/B3'],'actualZIPChartRangesVerified':['$B$2:$B$3','$C$2:$C$3'],'officeEngineRenderOrRecalculationNotClaimed':True},
 'actualIndependentMarks':[
  {'task':1,'components':[2,0,0],'reasons':['Konkrete Frage/durchführbarer Plan.','Gesamtprojekt-Ausführung nichtvorhanden; separates Worksheet ersetzt diese nicht.','Keine tatsächlich geleistete methodische Projekt-/Ergebnisauswertung.']},
  {'task':2,'components':[2,0,0],'reasons':['Rollen/Abhängigkeiten/Umplanung alsVorschlag.','Kein durchgeführter Team-/Rollendurchlauf.','Nur Toolabsicht, keine konkrete Evaluation.']},
  {'task':3,'components':[2,2,2],'reasons':['Tatsächlich erzeugte separate verknüpfte Charts mitrichtigen Daten/Bezügen.','Beide korrekten unabhängig geprüften Quotienten undRegistergrenze.','Begründete offene Erkundung samt Aggregat-/Personengrenze.']},
  {'task':4,'components':[2,2,2],'reasons':['Unveränderte vorherige ganze Präsentationsanalyse.','Zwei passende konkreteVerbesserungen.','Feedback-/Wirkungs-/Stereotypgrenzen.']}],
 'oldRawTotal':16,'missingWholePerformances':['Actually executed overall project cycle','Actually executed team/explicit role coordination'],'newWholeCap':14,'newFinalTotal':14,'newPass':False}
assert sum(sum(s['components']) for s in extra['actualIndependentMarks'])==16
save('actual-additional-complete-independent-sixteen-point-no-project-role-counteranswer-and-real-chart-XML-fixture.json',extra)

partial=[]
for g in new:
 partial.append({'examGoalId':g['id'],'case':'Actually executed required work with some inaccurate labels/arguments; presence is distinct from correctness.',
  'syntheticPredicateBoundaryNotObservedLearner':'All already required execution evidence present; partial marks assess mistakes, no fabricated learnerrecord.',
  'sampleRawStepMarks':[4,4,4,3],'actualRawTotal':15,'executionMissing':False,'wholeCapApplies':False,'finalTotal':15,'passesUnchanged15Threshold':True,
  'basis':'Both appended DE/EN conditions explicitly preserve ordinary partial credit for executed work with errors and disavow perfect6-point or100%requirements.'})
save('actual-four-old-whole-counteranswers-plus-one-new-case-and-two-partial-work-boundary-rechecks.json',{'schemaVersion':1,'wholePriorCounteranswersReused':rechecks,'newIndependentExtraCase':extra['id'],'partialExecutedWorkBoundaries':partial,'notLearnerMasteryOrNativeGraderHostClaim':True})
save('actual-ten-field-bounded-delta-and-two-individual-whole-successor-KEEP-judgments.json',{'schemaVersion':1,'changes':changes,'judgments':judgments,'allPriorWholeNumericFactsAndRubricsExceptStatedAppendsExact':True,'same24And15':True})
released=copy.deepcopy(new)
for g in released:g['examData']['reviewStatus']='released'
schema=json.loads((REPO/'docs/landscape-runtime.schema.json').read_text());small={'$ref':'#/$defs/goal','$defs':schema['$defs']};errors=[]
for g in released:
 errors.extend({'goalId':g['id'],'path':list(e.path),'message':e.message} for e in jsonschema.Draft202012Validator(small).iter_errors(g))
assert not errors
save('whole-two-current-successors.only-independent-machine-material-status-released.inert.json',released)
save('actual-two-goal-schema-and-current-four-counteranswers-plus-new-sixteen-case.delta-checks.json',{'schemaVersion':1,'goalSchemaErrors':errors,'independentWholeOldCounteranswers':len(rechecks),'additionalCompleteNewCounteranswer':1,'partialExecutedWorkBoundaries':2,'allAdversarialCasesRemainBelow15':True,'twoNewMachineMaterialKEEPNotStrictPromotion':2,'dependencyLocationOutsideCurricula':'/tmp/skillpilot-e40-independent-counterworkbook-deps','noNewRuntimeDependency':True})
print(json.dumps({'twoSuccessorKEEP':2,'oldCounteranswersRechecked':len(rechecks),'oldTotals':[r['oldActualRawTotal'] for r in rechecks],'newTotals':[r['newTotal'] for r in rechecks],'extraOld16ToNew14':True,'sameMaximum24AndPassing15':True,'goalSchemaErrors':errors},indent=2))
